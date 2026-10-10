# tests/test_report_command.py
"""The report command (reports spec 9)."""
from __future__ import annotations

import datetime
import hashlib
import io
import json
import sys

import pytest

pytest.importorskip("docx", reason="the report extras are not installed: pip install -e '.[report]'")

import leed_daylight.drivers.report as report_module
import leed_daylight.drivers.run as run_module
from deliver_folder import CASE, Option1Fake, build_deliver_project, edit_results
from option2_folder import load_chain
from stub_program import make_program

TODAY = datetime.date(2030, 1, 15)


@pytest.fixture
def folder(tmp_path, write_example, write_wea):
    return build_deliver_project(tmp_path, write_example, write_wea)


def _run(folder, **kwargs):
    out = io.StringIO()
    path = report_module.report_folder(folder, renderer="pillow", out=out, today=TODAY, **kwargs)
    return path, out.getvalue()


def test_every_report_is_written(folder):
    path, printed = _run(folder)

    assert path == folder / "_run" / "reports" / CASE
    for option in ("option1", "option2", "views"):
        assert (path / option / f"00_PROJECT_{option}.docx").is_file()
        assert (path / option / f"L0_{option}.docx").is_file()
        assert any((path / option / "plots").glob("*.png"))
    assert (path / "option1" / "glare_narrative.docx").is_file()
    assert (folder / "glare_narrative.md").is_file()
    assert "00_PROJECT_option1.docx" in printed
    assert "open items in glare_narrative.md" in printed


def test_the_manifest_ties_each_report_to_its_results(folder):
    path, _ = _run(folder)

    manifest = json.loads((path / report_module.REPORT_MANIFEST).read_text())
    assert manifest["case"] == CASE
    assert manifest["created"] == "2030-01-15"
    assert manifest["renderer"] == "pillow"
    results = folder / "_run" / "option2" / CASE / "option2_results.json"
    points = results.parent / "points.csv"
    assert manifest["options"]["option2"]["results_sha256"] == hashlib.sha256(results.read_bytes()).hexdigest()
    assert manifest["options"]["option2"]["points_sha256"] == hashlib.sha256(points.read_bytes()).hexdigest()
    assert "option2/00_PROJECT_option2.docx" in manifest["options"]["option2"]["files"]
    assert manifest["skipped"] == {}


def test_a_missing_option_is_skipped_and_printed(tmp_path, write_example, write_wea):
    folder = build_deliver_project(tmp_path, write_example, write_wea, options=("option2",))

    path, printed = _run(folder)

    assert not (path / "option1").exists()
    assert "skipped option1" in printed
    assert not (folder / "glare_narrative.md").exists()


def test_a_rerun_removes_stale_report_files(folder):
    path, _ = _run(folder)
    leftover = path / "option2" / "old_storey_option2.docx"
    leftover.write_text("stale", newline="\n")

    _run(folder)

    assert not leftover.exists()


def test_one_option_can_be_asked_for(folder):
    path, _ = _run(folder, options=("views",))

    assert (path / "views" / "00_PROJECT_views.docx").is_file()
    assert not (path / "option1").exists()


def test_main_reports_a_refusal_in_one_line(tmp_path, write_example, write_wea, capsys):
    folder = build_deliver_project(tmp_path, write_example, write_wea, options=())

    code = report_module.main([str(folder), "--renderer", "pillow"])

    assert code == 1
    err = capsys.readouterr().err
    assert err.startswith("error: no usable results")
    assert "Traceback" not in err


def test_main_refuses_an_unknown_renderer(folder, capsys):
    with pytest.raises(SystemExit):
        report_module.main([str(folder), "--renderer", "gnuplot"])


# --- Pre-flight ruling F1: out=None is resolved to sys.stdout at call time,
# so main's own default (no out= passed) is still captured by capsys.

def test_main_prints_to_the_current_stdout(folder, capsys):
    code = report_module.main([str(folder), "--renderer", "pillow"])

    assert code == 0
    out = capsys.readouterr().out
    assert "00_PROJECT_option1.docx" in out


# --- Pre-flight ruling F3: the manifest is merged, not rewritten from
# scratch, so a narrower rerun (--option views) keeps the other options'
# entries -- Task 12's package step depends on them.

def test_the_manifest_survives_a_narrower_rerun(folder):
    path, _ = _run(folder)

    _run(folder, options=("views",))

    manifest = json.loads((path / report_module.REPORT_MANIFEST).read_text())
    assert "option1" in manifest["options"]
    assert "option2" in manifest["options"]
    assert "views" in manifest["options"]
    # the option1/option2 files this narrower run did not touch are untouched
    assert (path / "option1" / "00_PROJECT_option1.docx").is_file()


def test_a_narrower_rerun_can_replace_a_skip_with_a_result(tmp_path, write_example, write_wea):
    # option1 has no results yet: the first full report skips it.
    folder = build_deliver_project(tmp_path, write_example, write_wea, options=("option2", "views"))
    path, _ = _run(folder)
    manifest = json.loads((path / report_module.REPORT_MANIFEST).read_text())
    assert "option1" in manifest["skipped"]
    assert "option1" not in manifest["options"]
    option2_files = manifest["options"]["option2"]["files"]
    views_files = manifest["options"]["views"]["files"]

    # Now option1 becomes usable, and a narrower `--option option1` rerun
    # (options=("option1",)) is asked for.
    _, prepared, _, _, _ = load_chain(folder)
    run_module.run_folder(folder, CASE, engine=Option1Fake(prepared), out=io.StringIO())
    _run(folder, options=("option1",))

    manifest = json.loads((path / report_module.REPORT_MANIFEST).read_text())
    assert "option1" not in manifest["skipped"]
    assert "option1" in manifest["options"]
    assert "option1/00_PROJECT_option1.docx" in manifest["options"]["option1"]["files"]
    # option2/views were not touched by this narrower rerun -- F3.
    assert manifest["options"]["option2"]["files"] == option2_files
    assert manifest["options"]["views"]["files"] == views_files


# --- Pre-flight ruling F4: python-docx-dependent modules are imported
# inside report_folder, in a try; a missing docx is a named ReportError
# naming python-docx and the install command, never a traceback.

def _block_docx(monkeypatch):
    """Simulate python-docx not being installed: forget every cached
    ``docx``/``docx.*`` module and every module of ours that imports it, and
    make a fresh ``import docx`` fail. Blocking only the top-level ``docx``
    name is not enough -- ``docx.oxml``, ``docx.shared`` and so on are
    resolved straight from ``sys.modules`` by their own full dotted name
    once cached, without consulting ``docx`` itself, so every module that
    was already loaded by an earlier test would otherwise still work."""
    for name in list(sys.modules):
        if name == "docx" or name.startswith("docx.") or name.startswith("leed_daylight.deliver.word") \
                or name == "leed_daylight.deliver.narrative":
            monkeypatch.delitem(sys.modules, name, raising=False)
    monkeypatch.setitem(sys.modules, "docx", None)


def test_a_missing_docx_is_a_named_refusal(folder, monkeypatch):
    _block_docx(monkeypatch)

    with pytest.raises(report_module.ReportError, match="docx"):
        report_module.report_folder(folder, renderer="pillow")


def test_main_reports_a_missing_docx_without_a_traceback(folder, monkeypatch, capsys):
    _block_docx(monkeypatch)

    code = report_module.main([str(folder), "--renderer", "pillow"])

    assert code == 1
    err = capsys.readouterr().err
    assert err.startswith("error:")
    assert "docx" in err
    assert 'pip install -e ".[report]"' in err
    assert "Traceback" not in err


def test_main_reports_a_missing_renderer_library_without_a_traceback(folder, monkeypatch, capsys):
    import leed_daylight.deliver.figures.base as base_module

    real_import_module = base_module.importlib.import_module

    def fake_import_module(name, *args, **kwargs):
        if name.endswith(".mpl"):
            raise ImportError("No module named 'matplotlib'", name="matplotlib")
        return real_import_module(name, *args, **kwargs)

    monkeypatch.setattr(base_module.importlib, "import_module", fake_import_module)

    code = report_module.main([str(folder), "--renderer", "matplotlib"])

    assert code == 1
    err = capsys.readouterr().err
    assert err.startswith("error:")
    assert "matplotlib" in err
    assert "Traceback" not in err


# --- Pre-flight ruling F6: the manifest records the glare narrative
# draft's own sha256 and the case it was rendered for, so a package step
# can tell whether the narrative on disk still matches the report it packages.

def test_the_manifest_records_the_narrative_file(folder):
    path, _ = _run(folder)

    manifest = json.loads((path / report_module.REPORT_MANIFEST).read_text())
    md_path = folder / "glare_narrative.md"
    assert manifest["options"]["option1"]["narrative_sha256"] == hashlib.sha256(md_path.read_bytes()).hexdigest()
    assert manifest["options"]["option1"]["narrative_case"] == CASE


# --- Pre-flight ruling F16: every file written is printed, including plots.

def test_every_written_file_is_printed_including_plots(folder):
    _, printed = _run(folder)

    assert "storeys_sda.png" in printed
    assert "storeys_pass.png" in printed
    assert "storeys_quality.png" in printed


# --- data.py:283's tool_version mismatch warning is part of CaseData.warnings,
# which report_folder prints; confirm it reaches the command's own output.

def test_a_tool_version_mismatch_warning_is_printed(folder):
    edit_results(folder, "option1", lambda payload: payload.__setitem__("tool_version", "0.0.0-old"))

    _, printed = _run(folder)

    assert "warning: results.json: written by tool version '0.0.0-old'" in printed


# -- Final review I5: a broken file is named in one error line ---------------

@pytest.mark.parametrize("text", ["[]\n", '{"options": ', '{"options": {"views": []}}\n'])
def test_main_refuses_a_malformed_report_manifest_in_one_line(folder, capsys, text):
    path, _ = _run(folder)
    (path / "report_manifest.json").write_text(text, newline="\n")

    code = report_module.main([str(folder), "--renderer", "pillow"])

    assert code == 1
    err = capsys.readouterr().err
    assert err.startswith("error: ") and err.count("\n") == 1
    assert str(path / "report_manifest.json") in err


def test_main_names_a_truncated_results_file(folder, capsys):
    results = folder / "_run" / "option2" / CASE / "option2_results.json"
    results.write_text(results.read_text()[:40], newline="\n")

    code = report_module.main([str(folder), "--renderer", "pillow"])

    assert code == 1
    err = capsys.readouterr().err
    assert err.startswith("error: ") and str(results) in err


# -- Final review I6 and M8: two storeys, the default renderer, end to end --

def test_a_two_storey_project_reports_end_to_end_with_the_default_renderer(tmp_path, write_example, write_wea):
    """The full Option 1 report (intent and sDA by space present), one docx
    per storey for every option, storey scoping, the "(this report)" mark,
    and the default matplotlib renderer (M8) -- none of it faked."""
    from docx import Document

    from deliver_folder import EXCLUDED_ROOM, UPPER_ROOM, build_two_storey_project, doc_headings, doc_text, table_after
    from leed_daylight.deliver.word import common
    from leed_daylight.deliver.word.option1 import NO_SPACE_RULE_WARNING
    from option2_folder import ROOM

    pytest.importorskip("matplotlib")
    folder = build_two_storey_project(tmp_path, write_example, write_wea)

    path = report_module.report_folder(folder, out=io.StringIO(), today=TODAY)

    manifest = json.loads((path / report_module.REPORT_MANIFEST).read_text())
    assert manifest["renderer"] == "matplotlib"
    for option in ("option1", "option2", "views"):
        docx = sorted(p.name for p in (path / option).glob("*.docx") if p.name != "glare_narrative.docx")
        assert docx == [f"00_PROJECT_{option}.docx", f"L0_{option}.docx", f"L1_{option}.docx"], option
        assert f"{option}/L1_{option}.docx" in manifest["options"][option]["files"]

    project = Document(path / "option1" / "00_PROJECT_option1.docx")
    text = doc_text(project)
    assert "An invented daylight intent sentence." in text
    assert ("Heading 1", "sDA by space") in doc_headings(project)
    assert common.NO_INTENT_WARNING not in text and NO_SPACE_RULE_WARNING not in text
    assert EXCLUDED_ROOM not in text
    # "Spaces" appears at level 1 (the project) and level 3 (each storey).
    project_rooms = {row.cells[1].text for row in table_after(project, "Spaces", level=1).rows[1:-1]}
    assert project_rooms == {ROOM, UPPER_ROOM}
    storey_rooms = {row.cells[1].text for row in table_after(project, "Spaces", level=3).rows[1:-1]}
    assert storey_rooms == {ROOM}

    upper = Document(path / "option1" / "L1_option1.docx")
    by_storey = [row.cells[0].text for row in table_after(upper, "Results by storey").rows[1:]]
    assert by_storey == ["L0", "L1 (this report)"]
    upper_rooms = {row.cells[1].text for row in table_after(upper, "Spaces", level=1).rows[1:-1]}
    assert upper_rooms == {UPPER_ROOM}

    views = Document(path / "views" / "00_PROJECT_views.docx")
    assert "An invented views intent sentence." in doc_text(views)


def test_storey_labels_that_collide_as_file_names_are_refused(tmp_path, write_example, write_wea, capsys):
    """Final review M7: "L 1" and "L_1" both become L_1_<option>.docx."""
    from deliver_folder import TWO_STOREY_SPACES_CSV, build_two_storey_project

    folder = build_two_storey_project(tmp_path, write_example, write_wea, options=())
    (folder / "spaces.csv").write_text(
        TWO_STOREY_SPACES_CSV.replace("\nL0,", "\nL 1,").replace("\nL1,", "\nL_1,")
    , newline="\n")
    _, prepared, _, _, _ = load_chain(folder)
    run_module.run_folder(folder, CASE, engine=Option1Fake(prepared), out=io.StringIO())

    code = report_module.main([str(folder), "--renderer", "pillow"])

    assert code == 1
    err = capsys.readouterr().err
    assert "'L 1'" in err and "'L_1'" in err and "L_1" in err


def test_storey_labels_that_differ_only_in_case_are_refused(tmp_path, write_example, write_wea, capsys):
    """CLI spec 7.1: "L1" and "l1" are one file on a case-insensitive disk
    (the macOS default), so one storey report would overwrite the other."""
    from deliver_folder import TWO_STOREY_SPACES_CSV, build_two_storey_project

    folder = build_two_storey_project(tmp_path, write_example, write_wea, options=())
    (folder / "spaces.csv").write_text(
        TWO_STOREY_SPACES_CSV.replace("\nL0,Room North,", "\nl1,Room North,")
    , newline="\n")
    _, prepared, _, _, _ = load_chain(folder)
    run_module.run_folder(folder, CASE, engine=Option1Fake(prepared), out=io.StringIO())

    code = report_module.main([str(folder), "--renderer", "pillow"])

    assert code == 1
    err = capsys.readouterr().err
    assert "'l1'" in err and "'L1'" in err
    assert not (folder / "_run" / "reports" / CASE / "option1").exists()


# -- Atlas spec 11: report also builds the floor atlas -------------------------

#: A stand-in for Blender (never a real one): the version line, then a build that exits 1.
FAILING_BLENDER = (
    "import sys\n"
    "if sys.argv[1:2] == ['--version']:\n"
    "    print('Blender 5.1.1'); sys.exit(0)\n"
    "print('an invented failure')\n"
    "sys.exit(1)\n"
)
#: A stand-in for Blender that saves the file named by its last argument.
SAVING_BLENDER = (
    "import sys\n"
    "if sys.argv[1:2] == ['--version']:\n"
    "    print('Blender 5.1.1'); sys.exit(0)\n"
    "open(sys.argv[-1], 'wb').write(b'BLENDER-v501')\n"
)


def _blender_script(tmp_path, monkeypatch, body):
    exe = make_program(tmp_path / "bin", "blender", body)
    monkeypatch.setenv("BLENDER", str(exe))
    return exe


def _add_missing_drawing(folder):
    """An ``atlas:`` drawing whose image does not exist: the atlas refuses
    it, the Word reports do not read it."""
    import yaml

    raw = yaml.safe_load((folder / "project.yaml").read_text())
    raw["atlas"] = {"drawings": [{
        "storey": "L0", "image": "drawings/level0.png",
        "points": [{"pixel": [0, 0], "model": [0.0, 0.0]}, {"pixel": [100, 0], "model": [10.0, 0.0]}],
    }]}
    (folder / "project.yaml").write_text(yaml.safe_dump(raw, sort_keys=False), newline="\n")


def test_report_writes_the_atlas_and_its_manifest_entry(folder):
    path, printed = _run(folder)

    atlas = folder / "_run" / "atlas" / CASE
    assert sorted(p.name for p in atlas.iterdir()) == ["atlas.html", "bundle.npz"]
    lines = printed.splitlines()
    assert lines.index(f"wrote {atlas / 'atlas.html'}") > lines.index(
        f"wrote {path / 'views' / '00_PROJECT_views.docx'}")
    # Under the test guard BLENDER names a path that does not exist.
    assert any(line.startswith("Blender skipped, no model.blend: BLENDER is set to ") for line in lines)
    manifest = json.loads((path / report_module.REPORT_MANIFEST).read_text())
    assert manifest["atlas"]["files"] == ["atlas.html", "bundle.npz"]
    assert sorted(manifest["atlas"]["options"]) == ["option1", "option2", "views"]
    results = folder / "_run" / "views" / CASE / "views_results.json"
    assert manifest["atlas"]["options"]["views"]["results_sha256"] == hashlib.sha256(
        results.read_bytes()).hexdigest()


def test_an_existing_atlas_entry_is_replaced(folder):
    path, _ = _run(folder)
    manifest_path = path / report_module.REPORT_MANIFEST
    manifest = json.loads(manifest_path.read_text())
    manifest["atlas"] = {"files": ["old.html"], "inputs": {}, "options": {}}
    manifest_path.write_text(json.dumps(manifest), newline="\n")

    _run(folder)

    atlas = json.loads(manifest_path.read_text())["atlas"]
    assert atlas["files"] == ["atlas.html", "bundle.npz"]
    assert sorted(atlas["options"]) == ["option1", "option2", "views"]


def test_an_atlas_failure_keeps_the_word_entries_and_drops_the_atlas_entry(folder, capsys):
    path, _ = _run(folder)
    _add_missing_drawing(folder)

    code = report_module.main([str(folder), "--renderer", "pillow"])

    assert code == 1
    err = capsys.readouterr().err
    assert err.startswith("error: atlas.drawings[0] (storey 'L0', image 'drawings/level0.png'): image not found")
    assert err.count("\n") == 1
    manifest = json.loads((path / report_module.REPORT_MANIFEST).read_text())
    assert "atlas" not in manifest
    assert sorted(manifest["options"]) == ["option1", "option2", "views"]
    assert (path / "views" / "00_PROJECT_views.docx").is_file()


def test_a_blender_failure_is_a_warning(folder, tmp_path, monkeypatch, capsys):
    _blender_script(tmp_path, monkeypatch, FAILING_BLENDER)

    code = report_module.main([str(folder), "--renderer", "pillow"])

    assert code == 0
    printed = capsys.readouterr()
    log = folder / "_run" / "atlas" / CASE / "blender.log"
    assert (f"warning: Blender failed, no model.blend: Blender exited 1; its output is in {log}"
            in printed.out.splitlines())
    assert printed.err == ""
    manifest = json.loads((folder / "_run" / "reports" / CASE / report_module.REPORT_MANIFEST).read_text())
    assert manifest["atlas"]["files"] == ["atlas.html", "bundle.npz"]


def test_a_blender_that_saves_adds_model_blend_to_the_entry(folder, tmp_path, monkeypatch, capsys):
    exe = _blender_script(tmp_path, monkeypatch, SAVING_BLENDER)

    assert report_module.main([str(folder), "--renderer", "pillow"]) == 0

    out = capsys.readouterr().out
    assert f"model.blend built by Blender 5.1.1 ({exe})" in out
    manifest = json.loads((folder / "_run" / "reports" / CASE / report_module.REPORT_MANIFEST).read_text())
    assert manifest["atlas"]["files"] == ["atlas.html", "bundle.npz", "model.blend"]


def test_no_blender_is_passed_through(folder, tmp_path, monkeypatch, capsys):
    _blender_script(tmp_path, monkeypatch, SAVING_BLENDER)

    assert report_module.main([str(folder), "--renderer", "pillow", "--no-blender"]) == 0

    assert "Blender" not in capsys.readouterr().out
    assert sorted(p.name for p in (folder / "_run" / "atlas" / CASE).iterdir()) == ["atlas.html", "bundle.npz"]

    _, printed = _run(folder, blender=False)
    assert "Blender" not in printed


# -- Final review M1: a bad atlas entry is replaced, never a loop --------------

def test_a_bad_atlas_entry_is_replaced_by_the_report(folder, capsys):
    path, _ = _run(folder)
    manifest_path = path / report_module.REPORT_MANIFEST
    manifest = json.loads(manifest_path.read_text())
    manifest["atlas"] = {"files": ["../x"], "inputs": {}, "options": {}}
    manifest_path.write_text(json.dumps(manifest), newline="\n")

    assert report_module.main([str(folder), "--renderer", "pillow", "--no-blender"]) == 0

    assert capsys.readouterr().err == ""
    atlas = json.loads(manifest_path.read_text())["atlas"]
    assert atlas["files"] == ["atlas.html", "bundle.npz"]
    assert sorted(atlas["options"]) == ["option1", "option2", "views"]


# -- Final review M4: the atlas's skipped options are printed too --------------

def test_report_for_one_option_prints_why_the_atlas_left_others_out(tmp_path, write_example, write_wea):
    folder = build_deliver_project(tmp_path, write_example, write_wea, options=("option2", "views"))

    _, printed = _run(folder, options=("views",))

    lines = printed.splitlines()
    skipped = [line for line in lines if line.startswith("skipped option1: ")]
    assert len(skipped) == 1
    assert skipped[0].startswith(f"skipped option1: no option1 results for case {CASE!r}")
    assert not any(line.startswith(("skipped views", "skipped option2")) for line in lines)


def test_an_option_the_report_skipped_is_printed_once(tmp_path, write_example, write_wea):
    folder = build_deliver_project(tmp_path, write_example, write_wea, options=("option2", "views"))

    _, printed = _run(folder)

    assert len([line for line in printed.splitlines() if line.startswith("skipped option1: ")]) == 1
