# tests/deliver_folder.py
"""A fictional project folder with results for all three options, for the
deliver tests (reports plan Task 5).

Built on tests/option2_folder.py's invented one-room project (6 m x 4 m x
3 m, "Room North", storey L0). Option 1 and Option 2 run through their real
drivers with fake engines; Quality Views results are written by hand in the
driver's own format. Everywhere, the points west of x = 3 m pass and the
points east of it fail. Every number is invented and round.
"""
from __future__ import annotations

import csv
import io
import json
import pathlib

import numpy as np

import leed_daylight.drivers.option2 as option2_module
import leed_daylight.drivers.run as run_module
from leed_daylight.constants import TOOL_VERSION
from leed_daylight.project.config import load_project
from leed_daylight.project.inputs import input_fingerprints
from option2_folder import CREDIT_RULES, MATERIALS_CSV, ROOM, SCENE_RAD, build_option2_project, load_chain

CASE = "base_case"
OPTIONS = ("option1", "option2", "views")
N_HOURS = 8760
N_OCCUPIED = 3650  # project_example.yaml: 08:00-18:00 every day


def _dump(payload: dict) -> str:
    """The Global Constraints' JSON form (F17): sorted, indented, no NaN."""
    return json.dumps(payload, sort_keys=True, indent=2, allow_nan=False) + "\n"


SPACES_CSV = """\
floor,room,space_type,regularly_occupied,area_m2,centroid_x,centroid_y,level_z,view_preserving_shades,merge_into,include_base_case,views_base_case,reason
L0,Room North,office,yes,24,3.0,2.0,0.0,no,,1,1,
"""

VIEWS = {"eye_height_m": 1.2, "eye_height_source": "Fictional example: seated eye height"}


class Option1Fake:
    """Sky lux 400 west of x = 3 m and 100 east of it; no direct sun. Each
    window group sees only its own room's grid; the residual pass covers
    every grid once."""

    def __init__(self, prepared):
        self._by_room = {}
        for grid in prepared.grids:
            n = grid.points.shape[0]
            sky = np.where(grid.points[:, 0] < 3.0, 400.0, 100.0)[:, None] * np.ones((1, N_OCCUPIED))
            self._by_room[grid.room] = ((grid.room,) * n, sky)

    @property
    def n_hours(self):
        return N_HOURS

    def direct(self, group, hours):
        index, sky = self._by_room[group.room]
        return index, np.zeros_like(sky)

    def sky(self, group, hours, open_schedule):
        index, sky = self._by_room[group.room]
        return index, sky, None

    def residual(self, hours):
        return [(index, np.zeros_like(sky), np.zeros_like(sky)) for index, sky in self._by_room.values()]


class Option2Fake:
    """500 lux west of x = 3 m and 100 lux east of it, at every sky."""

    def __init__(self, prepared):
        points = np.vstack([g.points for g in prepared.grids])
        self._lux = np.where(points[:, 0] < 3.0, 500.0, 100.0)

    def illuminance(self, moment):
        return self._lux


def _write_views_results(folder: pathlib.Path, prepared) -> None:
    """Every grid of ``prepared``: the points west of x = 3 m have a view."""
    case_dir = folder / "_run" / "views" / CASE
    case_dir.mkdir(parents=True, exist_ok=True)
    spaces, rows = [], []
    for grid in prepared.grids:
        west = grid.points[:, 0] < 3.0
        share = 100.0 * float(west.sum()) / west.shape[0]
        spaces.append({
            "room": grid.room, "storey": grid.storey, "area_m2": 24.0, "points": int(west.shape[0]),
            "type2_pct": 100.0, "type3_pct": share, "quality_pct": share,
            "quality_area_m2": 24.0 * share / 100.0,
        })
        rows.extend(
            [grid.room, grid.storey, f"{x:.3f}", f"{y:.3f}", f"{z:.3f}",
             1, int(west[i]), int(west[i]), "12.00", 100, 0]
            for i, (x, y, z) in enumerate(grid.points)
        )
    total = sum(s["area_m2"] for s in spaces)
    quality = sum(s["quality_area_m2"] for s in spaces)
    share = 100.0 * quality / total
    payload = {
        "tool_version": TOOL_VERSION, "radiance_version": None, "case": CASE,
        "settings": {
            "eye_height_m": 1.2, "eye_height_source": VIEWS["eye_height_source"],
            "max_interior_glass": 1, "angular_step_deg": 2.0, "rays_per_point": 5400,
            "ground_z_m": 0.0, "moveable_furniture": "exclude", "rules": {"object_distance_m": 7.5},
        },
        "method_notes": ["Invented method note."],
        "quality_area_pct": share, "quality_point_pct": share,
        "type2_area_pct": 100.0, "type3_area_pct": share,
        "total_area_m2": total, "quality_area_m2": quality,
        "score": {"points": None, "band": None, "reason": "below the 75 % threshold"},
        "spaces": spaces,
        "excluded_spaces": [], "warnings": [],
        "inputs": input_fingerprints(load_project(folder / "project.yaml"), folder),
    }
    (case_dir / "views_results.json").write_text(_dump(payload), newline="\n")
    with open(case_dir / "points.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["room", "storey", "x", "y", "z", "type2", "type3", "quality",
                         "median_outside_m", "reaching_rays", "limit_rays"])
        writer.writerows(rows)


def build_deliver_project(tmp_path, write_example, write_wea, *, options=OPTIONS, mutate_more=None,
                          extra_files=None) -> pathlib.Path:
    def mutate(d):
        d["views"] = dict(VIEWS)
        if mutate_more is not None:
            mutate_more(d)

    folder = build_option2_project(tmp_path, write_example, write_wea, spaces_csv=SPACES_CSV, mutate=mutate)
    for name, text in (extra_files or {}).items():
        (folder / name).write_text(text, encoding="utf-8", newline="\n")
    _, prepared, _, _, _ = load_chain(folder)
    if "option1" in options:
        run_module.run_folder(folder, CASE, engine=Option1Fake(prepared), out=io.StringIO())
    if "option2" in options:
        option2_module.option2_folder(folder, CASE, engine=Option2Fake(prepared), out=io.StringIO())
    if "views" in options:
        _write_views_results(folder, prepared)
    return folder


#: The two-storey variant (final review I6): Room North on L0 as above, a
#: store room beside it on L0 that no case includes, and Room Upper on L1
#: (the same 6 m x 4 m room, 4 m higher, with its own south window).
UPPER_ROOM = "Room Upper"
EXCLUDED_ROOM = "Store East"

TWO_STOREY_SCENE_RAD = SCENE_RAD + """
floor_carpet polygon floor_store_east
0
0
12  8 0 0   12 0 0   12 4 0   8 4 0

floor_carpet polygon floor_room_upper
0
0
12  0 0 4   6 0 4   6 4 4   0 4 4

ceiling_tile polygon ceiling_room_upper
0
0
12  0 0 7   0 4 7   6 4 7   6 0 7

wall_paint polygon wall_room_upper_south
0
0
12  0 0 4   6 0 4   6 0 7   0 0 7

glass_clear polygon window_upper
0
0
12  2 0 5   4 0 5   4 0 6   2 0 6
"""

TWO_STOREY_SPACES_CSV = """\
floor,room,space_type,regularly_occupied,area_m2,centroid_x,centroid_y,level_z,view_preserving_shades,merge_into,include_base_case,views_base_case,reason
L0,Room North,office,yes,24,3.0,2.0,0.0,no,,1,1,
L0,Store East,storage,no,16,10.0,2.0,0.0,no,,0,0,not regularly occupied
L1,Room Upper,office,yes,24,3.0,2.0,4.0,no,,1,1,
"""

#: The credit rules with the optional report keys present: the intent
#: sentences and the per-space sDA limits (invented wording and numbers).
TWO_STOREY_CREDIT_RULES = CREDIT_RULES.replace(
    'version: "Example rules v0"\n',
    'version: "Example rules v0"\nintent: "An invented daylight intent sentence."\n',
).replace(
    "    per_space_narrative_trigger_pct: 10.0\n",
    "    per_space_narrative_trigger_pct: 10.0\n  space_rule_sda_pct: [55.0, 75.0]\n",
) + """\
views:
  intent: "An invented views intent sentence."
  version: "Example views rules v0"
  quality_definition: "type2_and_type3"
  threshold_pct: 75.0
  points: 1
  vlt_min_exclusive: 0.40
  head_height_multiple: 3.0
  object_distance_m: 7.5
  obstruction_exception_m: 0.30
  method: {distance_statistic: "median", min_outside_distance_m: 2.0, panorama_vertical_deg: 60.0}
"""


def build_two_storey_project(tmp_path, write_example, write_wea, *, options=OPTIONS) -> pathlib.Path:
    """The two-storey variant of :func:`build_deliver_project` (I6)."""
    def mutate(d):
        d["views"] = dict(VIEWS)

    folder = build_option2_project(tmp_path, write_example, write_wea, spaces_csv=TWO_STOREY_SPACES_CSV,
                                   mutate=mutate, scene_rad=TWO_STOREY_SCENE_RAD)
    (folder / "credit_rules.yaml").write_text(TWO_STOREY_CREDIT_RULES, newline="\n")
    _, prepared, _, _, _ = load_chain(folder)
    if "option1" in options:
        run_module.run_folder(folder, CASE, engine=Option1Fake(prepared), out=io.StringIO())
    if "option2" in options:
        option2_module.option2_folder(folder, CASE, engine=Option2Fake(prepared), out=io.StringIO())
    if "views" in options:
        _write_views_results(folder, prepared)
    return folder


def doc_headings(doc) -> list[tuple[str, str]]:
    """The (style name, text) of every heading/title paragraph in ``doc``
    (reports Task 6+ word tests; F11 -- one shared helper, not a private
    copy in every ``test_deliver_word_*.py`` file)."""
    return [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith(("Heading", "Title"))]


def doc_text(doc) -> str:
    """Every paragraph's text, then every table cell's text, one per line --
    what the word report tests search for a phrase in (F11, see doc_headings)."""
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def table_after(doc, heading_text: str, level: int | None = None):
    """The first table in document order that follows the heading paragraph
    with this text (reports Task 7+; F11 -- one shared helper, not a
    private copy in every ``test_deliver_word_*.py`` file).

    A document can carry the same heading text at more than one level (say,
    a project-wide "Spaces" table at level 1 and a storey-scoped one at
    level 3, inside a "Storey <label>" subsection); pass ``level`` (0 for
    the title, otherwise the heading level) to pick the right one. The
    search matches only the FIRST heading paragraph with this text (and
    level, when given), then returns the first ``w:tbl`` element that
    follows it -- it never resumes matching headings afterwards, so a
    repeated heading further down the document is never mistaken for the
    one asked for.
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    style_name = None if level is None else ("Title" if level == 0 else f"Heading {level}")
    found = False
    for child in doc.element.body.iterchildren():
        if not found and child.tag == qn("w:p"):
            paragraph = Paragraph(child, doc)
            if paragraph.text == heading_text and (style_name is None or paragraph.style.name == style_name):
                found = True
        elif found and child.tag == qn("w:tbl"):
            return Table(child, doc)
    where = "" if level is None else f" (level {level})"
    raise AssertionError(f"no table found after heading {heading_text!r}{where}")


def edit_results(folder: pathlib.Path, option: str, change) -> None:
    """Load one option's results JSON, pass it through ``change`` (edits in place), write it back."""
    name = {"option1": "results.json", "option2": "option2_results.json", "views": "views_results.json"}[option]
    path = folder / "_run" / option / CASE / name
    payload = json.loads(path.read_text())
    change(payload)
    path.write_text(_dump(payload), newline="\n")


#: The roof-glass variant (skylights plan, Tasks 11-14): the same invented
#: 6 m x 4 m x 3 m "Room North", its south window in a wall with a real
#: opening, a ceiling at z = 3.0 with two 1 m x 1 m holes, a clear skylight
#: over one hole and a diffusing panel over the other, both flat at z = 3.4
#: (the tool builds a 0.40 m well from each down to the ceiling). Every
#: number is invented and round.
ROOF_ROOM = ROOM

_ROOF_MATERIALS = """\
void plastic wall_paint
0
0
5  0.6 0.6 0.6  0 0

void plastic ceiling_tile
0
0
5  0.8 0.8 0.8  0 0

void plastic floor_carpet
0
0
5  0.3 0.3 0.3  0 0

void glass glass_clear
0
0
3  0.76 0.76 0.76

void glass glass_roof
0
0
3  0.76 0.76 0.76

void trans diffuser_panel
0
0
7  0.8 0.8 0.8  0 0  0.5  0
"""

_ROOF_ROOM_SURFACES = """
floor_carpet polygon floor_room_north
0
0
12  0 0 0   6 0 0   6 4 0   0 4 0

ceiling_tile polygon ceiling_north_south_strip
0
0
12  0 0 3   0 1.5 3   6 1.5 3   6 0 3

ceiling_tile polygon ceiling_north_north_strip
0
0
12  0 2.5 3   0 4 3   6 4 3   6 2.5 3

ceiling_tile polygon ceiling_north_west_piece
0
0
12  0 1.5 3   0 2.5 3   1 2.5 3   1 1.5 3

ceiling_tile polygon ceiling_north_middle_piece
0
0
12  2 1.5 3   2 2.5 3   4 2.5 3   4 1.5 3

ceiling_tile polygon ceiling_north_east_piece
0
0
12  5 1.5 3   5 2.5 3   6 2.5 3   6 1.5 3

wall_paint polygon wall_south_below
0
0
12  0 0 0   6 0 0   6 0 1   0 0 1

wall_paint polygon wall_south_above
0
0
12  0 0 2   6 0 2   6 0 3   0 0 3

wall_paint polygon wall_south_west
0
0
12  0 0 1   2 0 1   2 0 2   0 0 2

wall_paint polygon wall_south_east
0
0
12  4 0 1   6 0 1   6 0 2   4 0 2
"""

_ROOF_WINDOW = """
glass_clear polygon window_ext
0
0
12  2 0 1   4 0 1   4 0 2   2 0 2
"""

_ROOF_GLASS = """
glass_roof polygon skylight_a
0
0
12  1 1.5 3.4   2 1.5 3.4   2 2.5 3.4   1 2.5 3.4

diffuser_panel polygon diffuser_a
0
0
12  4 1.5 3.4   5 1.5 3.4   5 2.5 3.4   4 2.5 3.4
"""

ROOF_SCENE_RAD = _ROOF_MATERIALS + _ROOF_ROOM_SURFACES + _ROOF_WINDOW + _ROOF_GLASS
#: One clear skylight and one diffusing panel, no window (the smoke run).
SKYLIGHT_ONLY_SCENE_RAD = _ROOF_MATERIALS + _ROOF_ROOM_SURFACES + _ROOF_GLASS

#: The smoke run's own scene: a closed room (four solid walls, no window), so
#: the only way in for the sun is through the roof glass. The clear skylight
#: sits 0.5 m further south than in the other roof scenes, over a ceiling hole
#: at x 1-2, y 1-2, so that some lit suns of the invented weather reach the
#: back row of points through it (the smoke run's control). The diffusing
#: panel and its hole are unchanged. Every number is invented and round.
_CLOSED_ROOM_SURFACES = """
floor_carpet polygon floor_room_north
0
0
12  0 0 0   6 0 0   6 4 0   0 4 0

ceiling_tile polygon ceiling_closed_south_strip
0
0
12  0 0 3   0 1 3   6 1 3   6 0 3

ceiling_tile polygon ceiling_closed_north_strip
0
0
12  0 2.5 3   0 4 3   6 4 3   6 2.5 3

ceiling_tile polygon ceiling_closed_west_piece
0
0
12  0 1 3   0 2.5 3   1 2.5 3   1 1 3

ceiling_tile polygon ceiling_closed_behind_skylight
0
0
12  1 2 3   1 2.5 3   2 2.5 3   2 2 3

ceiling_tile polygon ceiling_closed_middle_piece
0
0
12  2 1 3   2 2.5 3   4 2.5 3   4 1 3

ceiling_tile polygon ceiling_closed_before_diffuser
0
0
12  4 1 3   4 1.5 3   5 1.5 3   5 1 3

ceiling_tile polygon ceiling_closed_east_piece
0
0
12  5 1 3   5 2.5 3   6 2.5 3   6 1 3

wall_paint polygon wall_closed_south
0
0
12  0 0 0   6 0 0   6 0 3   0 0 3

wall_paint polygon wall_closed_north
0
0
12  0 4 0   0 4 3   6 4 3   6 4 0

wall_paint polygon wall_closed_east
0
0
12  6 0 0   6 4 0   6 4 3   6 0 3

wall_paint polygon wall_closed_west
0
0
12  0 0 0   0 0 3   0 4 3   0 4 0
"""

_CLOSED_ROOM_GLASS = """
glass_roof polygon skylight_a
0
0
12  1 1 3.4   2 1 3.4   2 2 3.4   1 2 3.4

diffuser_panel polygon diffuser_a
0
0
12  4 1.5 3.4   5 1.5 3.4   5 2.5 3.4   4 2.5 3.4
"""

#: One clear skylight and one diffusing panel in a closed room (the smoke run).
CLOSED_SKYLIGHT_SCENE_RAD = _ROOF_MATERIALS + _CLOSED_ROOM_SURFACES + _CLOSED_ROOM_GLASS

ROOF_MATERIALS_CSV = MATERIALS_CSV + "glass_roof,glass,glazing,\ndiffuser_panel,trans,glazing_diffusing,\n"


def build_roof_glass_project(tmp_path, write_example, write_wea, *, options=("option1",),
                             scene_rad: str = ROOF_SCENE_RAD, ambient_full: str | None = None) -> pathlib.Path:
    """The roof-glass variant of :func:`build_deliver_project`. Only Option 1
    is run (with the fake engine), and only when asked for."""
    def mutate(d):
        d["views"] = dict(VIEWS)
        d["glazing"]["dirt_factor"] = {"vertical": 0.95, "skylight": 0.85}
        d["glazing"]["frame_share"]["skylight"] = 0.10
        d["glazing"]["skylight_well"] = 0.40
        d["glazing"]["skylight_well_reflectance"] = 0.70
        if ambient_full is not None:
            d["radiance"]["ambient"]["full"] = ambient_full

    folder = build_option2_project(tmp_path, write_example, write_wea, spaces_csv=SPACES_CSV,
                                   mutate=mutate, scene_rad=scene_rad)
    (folder / "materials.csv").write_text(ROOF_MATERIALS_CSV, newline="\n")
    if "option1" in options:
        _, prepared, _, _, _ = load_chain(folder)
        run_module.run_folder(folder, CASE, engine=Option1Fake(prepared), out=io.StringIO())
    return folder
