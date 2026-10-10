# tests/test_deliver_narrative.py
"""The glare narrative draft (reports spec 7)."""
from __future__ import annotations

import dataclasses
import types

import pytest

pytest.importorskip("docx", reason="the report extras are not installed: pip install -e '.[report]'")
from docx import Document

from deliver_folder import ROOM, build_deliver_project, build_roof_glass_project
from leed_daylight.deliver import narrative
from leed_daylight.deliver.data import load_case


@pytest.fixture
def case(tmp_path, write_example, write_wea):
    """The invented case, with its one space over the ASE trigger."""
    case = load_case(build_deliver_project(tmp_path, write_example, write_wea, options=("option1",)))
    results = dict(case.options["option1"].results)
    results["spaces"] = [dict(s, ase_pct=25.0, needs_glare_narrative=True) for s in results["spaces"]]
    results["narrative_spaces"] = [ROOM]
    option = dataclasses.replace(case.options["option1"], results=types.MappingProxyType(results))
    return dataclasses.replace(case, options=types.MappingProxyType({"option1": option}))


def test_first_run_writes_the_template_with_both_blocks(case, tmp_path):
    path = tmp_path / "glare_narrative.md"

    warnings = narrative.update_narrative(path, case)

    text = path.read_text()
    assert warnings == []
    assert text.startswith("# Glare narrative")
    assert "<!-- auto:results -->" in text and "<!-- /auto:spaces -->" in text
    assert f"| L0 | {ROOM} | 24.0 | 25.0 % |" in text
    assert narrative.count_marks(text) >= 3


def test_a_rerun_keeps_hand_text_and_filled_measures(case, tmp_path):
    path = tmp_path / "glare_narrative.md"
    narrative.update_narrative(path, case)
    text = path.read_text()
    text = text.replace("[ADD a short description of the glare strategy", "Our own glare strategy text. [ADD more")
    row_start = text.index(f"| L0 | {ROOM} |")
    row_end = text.index("\n", row_start)
    row = text[row_start:row_end]
    text = text.replace(row, row.rsplit("|", 2)[0] + "| External fins on the south facade |")
    path.write_text(text, newline="\n")

    narrative.update_narrative(path, case)

    again = path.read_text()
    assert "Our own glare strategy text." in again
    assert "External fins on the south facade" in again
    assert again.count("<!-- auto:spaces -->") == 1


def test_missing_markers_append_the_block_and_warn(case, tmp_path):
    path = tmp_path / "glare_narrative.md"
    path.write_text("# Glare narrative\n\nOnly my own text.\n", newline="\n")

    warnings = narrative.update_narrative(path, case)

    text = path.read_text()
    assert "Only my own text." in text
    assert "<!-- auto:results -->" in text and "<!-- auto:spaces -->" in text
    assert len(warnings) == 2
    assert "auto:results" in warnings[0]


def test_an_opener_without_its_closer_is_left_alone_and_warned(case, tmp_path):
    """Fix round 1, CRITICAL finding: an orphaned opener must never be
    silently treated as a missing block on a later run (that used to let a
    non-greedy regex match it to the *other* block's closer and overwrite
    the user's filled-in row with no warning)."""
    path = tmp_path / "glare_narrative.md"
    narrative.update_narrative(path, case)
    text = path.read_text().replace("<!-- /auto:spaces -->\n", "")
    path.write_text(text, newline="\n")
    corrupted = path.read_text()

    warnings1 = narrative.update_narrative(path, case)
    after1 = path.read_text()
    warnings2 = narrative.update_narrative(path, case)
    after2 = path.read_text()

    assert after1 == corrupted
    assert after2 == corrupted
    assert after1.count("<!-- auto:spaces -->") == 1
    assert len(warnings1) == 1 and len(warnings2) == 1
    assert "auto:spaces" in warnings1[0] and "malformed" in warnings1[0]
    assert "auto:spaces" in warnings2[0] and "malformed" in warnings2[0]


def test_a_closer_without_its_opener_is_left_alone_and_warned(case, tmp_path):
    path = tmp_path / "glare_narrative.md"
    narrative.update_narrative(path, case)
    text = path.read_text().replace("<!-- auto:spaces -->\n", "")
    path.write_text(text, newline="\n")
    corrupted = path.read_text()

    warnings1 = narrative.update_narrative(path, case)
    after1 = path.read_text()
    warnings2 = narrative.update_narrative(path, case)
    after2 = path.read_text()

    assert after1 == corrupted
    assert after2 == corrupted
    assert after1.count("<!-- /auto:spaces -->") == 1
    assert len(warnings1) == 1 and len(warnings2) == 1
    assert "auto:spaces" in warnings1[0] and "malformed" in warnings1[0]


def test_duplicated_openers_are_left_alone_and_warned(case, tmp_path):
    path = tmp_path / "glare_narrative.md"
    narrative.update_narrative(path, case)
    text = path.read_text().replace(
        "<!-- auto:spaces -->\n", "<!-- auto:spaces -->\n<!-- auto:spaces -->\n", 1
    )
    path.write_text(text, newline="\n")
    corrupted = path.read_text()

    warnings = narrative.update_narrative(path, case)
    after = path.read_text()

    assert after == corrupted
    assert after.count("<!-- auto:spaces -->") == 2
    assert len(warnings) == 1
    assert "auto:spaces" in warnings[0] and "malformed" in warnings[0]


def test_a_glare_measures_answer_with_a_literal_pipe_survives_a_rerun(case, tmp_path):
    """Fix round 1, minor finding: the naive split used to lose an answer
    that itself contained "|" instead of keeping it verbatim."""
    path = tmp_path / "glare_narrative.md"
    narrative.update_narrative(path, case)
    text = path.read_text()
    row_start = text.index(f"| L0 | {ROOM} |")
    row_end = text.index("\n", row_start)
    row = text[row_start:row_end]
    text = text.replace(row, row.rsplit("|", 2)[0] + "| Blinds | operated automatically |")
    path.write_text(text, newline="\n")

    narrative.update_narrative(path, case)

    again = path.read_text()
    assert "Blinds | operated automatically" in again


def test_render_docx_marks_the_open_items(case, tmp_path):
    path = tmp_path / "glare_narrative.md"
    narrative.update_narrative(path, case)
    out = tmp_path / "glare_narrative.docx"

    narrative.render_docx(path.read_text(), out, case)

    doc = Document(out)
    assert doc.paragraphs[0].text.startswith("Glare narrative")
    assert len(doc.tables) == 2  # the results table and the spaces table
    add_runs = [r for p in doc.paragraphs for r in p.runs if r.text.startswith("[ADD")]
    confirm_runs = [r for p in doc.paragraphs for r in p.runs if r.text.startswith("[CONFIRM")]
    assert add_runs and all(r.font.color.rgb is not None for r in add_runs)
    assert confirm_runs and all(r.bold and r.font.color.rgb is not None for r in confirm_runs)


def _with_narrative_spaces(case, names):
    """``case`` with its narrative list set to ``names`` (a design change, or
    another case whose spaces over the trigger differ)."""
    results = dict(case.options["option1"].results)
    results["narrative_spaces"] = list(names)
    option = dataclasses.replace(case.options["option1"], results=types.MappingProxyType(results))
    return dataclasses.replace(case, options=types.MappingProxyType({"option1": option}))


def _fill_measures(path, text_in):
    text = path.read_text()
    row_start = text.index(f"| L0 | {ROOM} |")
    row_end = text.index("\n", row_start)
    row = text[row_start:row_end]
    path.write_text(text.replace(row, row.rsplit("|", 2)[0] + f"| {text_in} |"), newline="\n")


@pytest.mark.parametrize("opener_suffix", [" ", "  \t"])
def test_an_opener_with_trailing_whitespace_is_still_replaced(case, tmp_path, opener_suffix):
    """Final review C1: an editor that leaves trailing spaces after an
    opener used to crash the report command ("substring not found")."""
    path = tmp_path / "glare_narrative.md"
    narrative.update_narrative(path, case)
    _fill_measures(path, "External fins on the south facade")
    text = path.read_text()
    for name in ("results", "spaces"):
        text = text.replace(f"<!-- auto:{name} -->\n", f"<!-- auto:{name} -->{opener_suffix}\n")
    path.write_text(text, newline="\n")

    warnings = narrative.update_narrative(path, case)

    again = path.read_text()
    assert warnings == []
    assert "External fins on the south facade" in again
    assert again.count("| Case |") == 1
    assert again.count(f"| L0 | {ROOM} |") == 1


def test_a_block_on_one_line_does_not_crash(case, tmp_path):
    """Final review C1: an opener and its closer on one line."""
    path = tmp_path / "glare_narrative.md"
    path.write_text(
        "# Glare narrative\n\nMy text.\n\n<!-- auto:results --><!-- /auto:results -->\n\n"
        "<!-- auto:spaces --> <!-- /auto:spaces -->\n"
    , newline="\n")

    warnings = narrative.update_narrative(path, case)

    again = path.read_text()
    assert warnings == []
    assert "My text." in again
    assert again.count("<!-- auto:results -->") == 1 and again.count("<!-- /auto:spaces -->") == 1
    assert "| Case |" in again and f"| L0 | {ROOM} |" in again
    assert narrative.update_narrative(path, case) == []
    assert path.read_text() == again


def test_a_filled_room_that_drops_off_the_trigger_list_is_kept_and_warned(case, tmp_path):
    """Final review C2 (Ruling FR-C2): filled Glare measures text is never
    lost when its room leaves the narrative list."""
    path = tmp_path / "glare_narrative.md"
    narrative.update_narrative(path, case)
    _fill_measures(path, "External fins on the south facade")

    warnings = narrative.update_narrative(path, _with_narrative_spaces(case, []))

    text = path.read_text()
    assert "External fins on the south facade" in text
    assert "Spaces no longer over the trigger" in text
    assert "No space has ASE over the trigger." in text
    assert len(warnings) == 1 and ROOM in warnings[0] and "no longer over the trigger" in warnings[0]


def test_a_room_that_returns_to_the_trigger_list_gets_its_text_back(case, tmp_path):
    path = tmp_path / "glare_narrative.md"
    narrative.update_narrative(path, case)
    _fill_measures(path, "External fins on the south facade")
    narrative.update_narrative(path, _with_narrative_spaces(case, []))

    warnings = narrative.update_narrative(path, case)

    text = path.read_text()
    assert warnings == []
    assert "Spaces no longer over the trigger" not in text
    row_start = text.index(f"| L0 | {ROOM} |")
    assert text[row_start:text.index("\n", row_start)].endswith("| External fins on the south facade |")


def test_an_unfilled_room_that_drops_off_the_list_is_not_carried(case, tmp_path):
    path = tmp_path / "glare_narrative.md"
    narrative.update_narrative(path, case)

    warnings = narrative.update_narrative(path, _with_narrative_spaces(case, []))

    text = path.read_text()
    assert warnings == []
    assert "Spaces no longer over the trigger" not in text
    assert ROOM not in text.split("<!-- auto:spaces -->")[1]


def test_a_pipe_in_a_room_name_keeps_the_table_and_the_measures(case, tmp_path):
    """Final review M9: a "|" in a room name is escaped in the table cell,
    so the row keeps its six cells and its filled text on a rerun."""
    room = "Office | East"
    results = dict(case.options["option1"].results)
    results["spaces"] = [dict(s, room=room) for s in results["spaces"]]
    results["narrative_spaces"] = [room]
    option = dataclasses.replace(case.options["option1"], results=types.MappingProxyType(results))
    storeys = tuple(dataclasses.replace(s, rooms=(*s.rooms, room)) for s in case.storeys)
    piped = dataclasses.replace(case, options=types.MappingProxyType({"option1": option}), storeys=storeys)
    path = tmp_path / "glare_narrative.md"
    narrative.update_narrative(path, piped)
    text = path.read_text()
    assert "| L0 | Office \\| East |" in text
    row_start = text.index("| L0 | Office \\| East |")
    row = text[row_start:text.index("\n", row_start)]
    path.write_text(text.replace(row, row.rsplit("|", 2)[0] + "| Blinds on the east facade |"), newline="\n")

    narrative.update_narrative(path, piped)

    again = path.read_text()
    assert "| L0 | Office \\| East |" in again and "Blinds on the east facade" in again
    out = tmp_path / "glare_narrative.docx"
    narrative.render_docx(again, out, piped)
    cells = [c.text for c in Document(out).tables[1].rows[1].cells]
    assert cells[1] == room and cells[5] == "Blinds on the east facade"


def _over_trigger(case):
    results = dict(case.options["option1"].results)
    results["spaces"] = [dict(s, ase_pct=25.0, needs_glare_narrative=True) for s in results["spaces"]]
    results["narrative_spaces"] = [ROOM]
    option = dataclasses.replace(case.options["option1"], results=types.MappingProxyType(results))
    return dataclasses.replace(case, options=types.MappingProxyType({"option1": option}))


def _spaces_block(text):
    return text[text.index("<!-- auto:spaces -->"):text.index("<!-- /auto:spaces -->")]


def test_diffusing_glass_in_a_space_over_the_trigger_is_named(tmp_path, write_example, write_wea):
    case = _over_trigger(load_case(build_roof_glass_project(tmp_path, write_example, write_wea)))
    path = tmp_path / "glare_narrative.md"

    narrative.update_narrative(path, case)
    narrative.update_narrative(path, case)

    block = _spaces_block(path.read_text())
    line = f"Diffusing glass in {ROOM} (Room_North_dif) needs no glare-control device"
    assert block.count(line) == 1
    assert "do not require glare-control devices" in block
    assert f"| L0 | {ROOM} | 24.0 | 25.0 % |" in block


def test_no_diffusing_line_under_the_trigger(tmp_path, write_example, write_wea):
    case = load_case(build_roof_glass_project(tmp_path, write_example, write_wea))
    assert not case.options["option1"].results.get("narrative_spaces")
    path = tmp_path / "glare_narrative.md"

    narrative.update_narrative(path, case)

    assert "Diffusing glass" not in path.read_text()


def test_no_diffusing_line_without_diffusing_glass(case, tmp_path):
    path = tmp_path / "glare_narrative.md"

    narrative.update_narrative(path, case)

    assert "Diffusing glass" not in path.read_text()
