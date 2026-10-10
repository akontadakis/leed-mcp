# leed_daylight/deliver/narrative.py
"""The glare narrative draft (reports spec 7).

The editable file is ``<project folder>/glare_narrative.md``. The first run
writes ``TEMPLATE``. Every run replaces only the text between the
``<!-- auto:NAME -->`` and ``<!-- /auto:NAME -->`` markers; everything else,
and the "Glare measures" text of each space, is the user's and is kept. A
block whose markers are missing (neither the opener nor the closer present)
is appended at the end, with a warning.

Fix round 1 (review finding, CRITICAL): a block whose markers are anything
*other* than "exactly one opener, one closer, opener before closer" or
"neither present" -- an opener with no closer, a closer with no opener,
duplicated markers, or a closer before its opener -- is malformed. Such a
block is left exactly as written (nothing appended, nothing replaced) and a
distinct warning names it and says to fix it by hand. The earlier version
scanned for well-formed blocks with a single regex and treated anything it
did not match as "missing", which on a second run silently matched the
orphaned opener to the *other* block's closer and overwrote the user's text
with no warning; ``_block_state`` below classifies a block before anything
is read or replaced, so a malformed block is never mistaken for a missing
one.

Fix round 1 (review finding, minor): a "Glare measures" answer that itself
contains a literal ``|`` used to be lost silently on the next run (the
naive ``str.split("|")`` produced more than six cells, so the row failed
the "is this a data row" check and was skipped). ``_measures`` now splits
each row with ``maxsplit=5``, so the sixth (last) cell keeps everything
after the five fixed leading cells, literal ``|`` included.

Final review C1: the body of a block runs from the end of the opener's
*line* (trailing spaces after the opener are fine) up to the closer, and it
is replaced by slicing, never by rebuilding the marker text and searching
for it. An opener and closer on one line is also handled: the new body goes
on its own lines between them.

Final review C2 (Ruling FR-C2): a space whose "Glare measures" text was
filled and which is no longer over the trigger (another case, or a design
change) is never dropped. Its row moves to a second table, "Spaces no longer
over the trigger", inside the same auto block, and a warning names it. When
the space is back over the trigger its text returns to the main table. Rows
still ``[ADD]`` are not carried.

Final review M9: a ``|`` in any generated cell (a room name, say) is written
as ``\\|`` so the row keeps its six cells; the parser and the docx renderer
split on unescaped ``|`` only and turn ``\\|`` back into ``|``.

The markdown-subset renderer (headings, paragraphs, bullets, pipe tables,
rules, bold, italic, code and the [ADD]/[CONFIRM] marks) is ported from
the author's earlier work (PORTED.md).
"""
from __future__ import annotations

import pathlib
import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..engine.jobs import refuse_link, replace_file
from ..engine.in_use import explain_in_use
from .data import CaseData
from .word import kit

TEMPLATE = """\
# Glare narrative, LEED v4.1 EQ Credit Daylight

## 1. Why this narrative is needed

LEED v4.1 asks for a narrative of how the design addresses glare in each regularly occupied space whose annual sunlight exposure (ASE) is over the per-space trigger. [CONFIRM that the spaces in section 4 are regularly occupied.]

## 2. Results

<!-- auto:results -->
<!-- /auto:results -->

## 3. Glare-control devices

[ADD the blinds or shades on each facade: product, operation (manual, or automatic with manual override), and the spaces they serve.]

[ADD a short description of the glare strategy for the spaces in section 4.]

## 4. Spaces with ASE over the trigger

<!-- auto:spaces -->
<!-- /auto:spaces -->
"""

_MARK = re.compile(r"\[(?:ADD|CONFIRM)\b[^\]]*\]")
_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+\*|`[^`]+`|\[(?:ADD|CONFIRM)\b[^\]]*\])", re.S)
_OPEN_ITEM = RGBColor(0xB0, 0x30, 0x20)  # dark red, so an unfilled mark is seen


def count_marks(text: str) -> int:
    return len(_MARK.findall(text))


def _results_block(case: CaseData) -> str:
    r = case.options["option1"].results
    trigger = float(case.rules["option1"]["ase"]["per_space_narrative_trigger_pct"])
    points = r["score"]["points"]
    return "\n".join([
        "| Result | Value |",
        "|---|---:|",
        f"| Case | {case.case} |",
        f"| sDA, floor-area weighted | {float(r['sda_area_pct']):.1f} % |",
        f"| ASE, floor-area weighted | {float(r['ase_area_pct']):.1f} % |",
        f"| Points | {points if points is not None else 'not awarded'} |",
        f"| Spaces with ASE over {trigger:g} % | {len(r.get('narrative_spaces', []))} |",
    ]) + "\n"


_PIPE = re.compile(r"(?<!\\)\|")


def _cell(value) -> str:
    """One generated table cell, with a literal ``|`` escaped (M9)."""
    return str(value).replace("|", "\\|")


def _split_row(line: str, maxsplit: int = 0) -> list[str]:
    """The cells of one markdown table row, split on unescaped ``|`` only."""
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    return [c.strip() for c in _PIPE.split(line, maxsplit)]


def _measures(body: str) -> dict[str, str]:
    kept = {}
    for line in body.splitlines():
        if not line.strip().startswith("|"):
            continue
        # maxsplit=5: the sixth cell (Glare measures) keeps everything after
        # the five fixed leading cells, including a literal "|" the user
        # typed in their answer (fix round 1, minor finding).
        cells = _split_row(line, 5)
        if len(cells) == 6 and cells[1] not in ("Space", "---") and not set(cells[1]) <= {"-", ":"}:
            kept[cells[1].replace("\\|", "|")] = cells[5]
    return kept


_SPACES_HEAD = ["| Storey | Space | Area m2 | ASE | sDA | Glare measures |", "|---|---|---:|---:|---:|---|"]

#: Skylights spec 7 (D9): the reason printed for diffusing glass in a space
#: over the trigger (LEED v4 BD+C Reference Guide, EQ Credit Daylight).
DIFFUSING_NO_DEVICE = (
    'the LEED v4 BD+C Reference Guide, EQ Credit Daylight, says "Diffused and translucent glazing systems '
    'do not require glare-control devices."'
)


def _is_filled(answer: str) -> bool:
    return answer.strip() not in ("", "[ADD]")


def _spaces_block(case: CaseData, kept: dict[str, str]) -> tuple[str, list[str]]:
    """The spaces block, and the rooms whose filled text was carried into
    the "no longer over the trigger" table (C2)."""
    r = case.options["option1"].results
    names = set(r.get("narrative_spaces", []))
    by_room = {s["room"]: s for s in r["spaces"]}

    def row(room: str, answer: str) -> str:
        s = by_room.get(room)
        if s is None:
            values = ["n/a", room, "n/a", "n/a", "n/a"]
        else:
            try:
                storey = case.storey_of(room)
            except ValueError:
                storey = "n/a"
            values = [storey, room, f"{float(s['area_m2']):.1f}", f"{float(s['ase_pct']):.1f} %",
                      f"{float(s['sda_pct']):.1f} %"]
        return "| " + " | ".join(_cell(v) for v in values) + f" | {answer} |"

    listed = [s["room"] for s in r["spaces"] if s["room"] in names]
    dropped = [room for room, answer in kept.items() if room not in listed and _is_filled(answer)]
    parts = []
    if listed:
        parts.append("\n".join([*_SPACES_HEAD, *(row(room, kept.get(room, "[ADD]")) for room in listed)]))
        diffusing: dict[str, list[str]] = {}
        for g in case.groups:
            if g.kind == "diffusing":
                diffusing.setdefault(g.room, []).append(g.group_id)
        parts.extend(
            f"Diffusing glass in {_cell(room)} ({', '.join(diffusing[room])}) needs no glare-control device: "
            + DIFFUSING_NO_DEVICE
            for room in listed if room in diffusing
        )
    else:
        parts.append("No space has ASE over the trigger.")
    if dropped:
        parts.append(f"**Spaces no longer over the trigger** in case {_cell(case.case)}; "
                     "their Glare measures text is kept here.")
        parts.append("\n".join([*_SPACES_HEAD, *(row(room, kept[room]) for room in dropped)]))
    return "\n\n".join(parts) + "\n", dropped


def _block_state(text: str, name: str) -> str:
    """Classify one auto block by counting its literal markers -- never by
    matching a single regex and treating a non-match as "missing" (that is
    what let a malformed block masquerade as a missing one; see the module
    docstring, fix round 1)."""
    open_tag, close_tag = f"<!-- auto:{name} -->", f"<!-- /auto:{name} -->"
    opens, closes = text.count(open_tag), text.count(close_tag)
    if opens == 0 and closes == 0:
        return "missing"
    if opens == 1 and closes == 1 and text.index(open_tag) < text.index(close_tag):
        return "well_formed"
    return "malformed"


def _body_span(text: str, name: str) -> tuple[int, int, str]:
    """``(start, end, lead)`` of a well-formed block's body: from the end of
    the opener's line (C1: trailing spaces after the opener are fine) to the
    closer. When the closer is on the opener's line, the body starts right
    after the opener and ``lead`` is a newline, so the new body goes on its
    own lines. Only called once ``_block_state`` has confirmed exactly one
    opener before exactly one closer."""
    open_tag, close_tag = f"<!-- auto:{name} -->", f"<!-- /auto:{name} -->"
    after_open = text.index(open_tag) + len(open_tag)
    end = text.index(close_tag, after_open)
    line_end = text.find("\n", after_open, end)
    if line_end == -1:
        return after_open, end, "\n"
    return line_end + 1, end, ""


def update_narrative(path: pathlib.Path, case: CaseData) -> list[str]:
    path = pathlib.Path(path)
    # Never read or written through a link (security review H1's class).
    refuse_link(path)
    text = path.read_text(encoding="utf-8") if path.is_file() else TEMPLATE
    names = ("results", "spaces")
    states = {name: _block_state(text, name) for name in names}

    kept = {}
    if states["spaces"] == "well_formed":
        start, end, _ = _body_span(text, "spaces")
        kept = _measures(text[start:end])
    spaces_body, dropped = _spaces_block(case, kept)
    fresh = {"results": _results_block(case), "spaces": spaces_body}

    warnings = []
    for name in names:
        state, body = states[name], fresh[name]
        open_tag, close_tag = f"<!-- auto:{name} -->", f"<!-- /auto:{name} -->"
        if state == "well_formed":
            start, end, lead = _body_span(text, name)
            text = text[:start] + lead + body + text[end:]
        elif state == "missing":
            text = text.rstrip("\n") + f"\n\n{open_tag}\n{body}{close_tag}\n"
            warnings.append(f"{path.name}: the auto:{name} markers were missing; the block was added at the end")
        else:
            warnings.append(
                f"{path.name}: the auto:{name} markers are malformed (an opener without its closer, a closer "
                "without its opener, duplicated markers, or a closer before its opener); that block was left "
                "exactly as written -- fix the markers by hand"
            )
    if dropped:
        warnings.append(
            f"{path.name}: {', '.join(dropped)} no longer over the trigger in case {case.case}; "
            "their Glare measures text is kept in the \"Spaces no longer over the trigger\" table"
        )
    replace_file(path.parent, path, text.encode("utf-8"))
    return warnings


def _runs(paragraph, text: str, size: int = 10) -> None:
    for piece in _INLINE.split(text):
        if not piece:
            continue
        plain = piece.strip("*`") if piece[0] in "*`" else piece
        run = paragraph.add_run(re.sub(r"\s+", " ", plain))
        run.font.size = Pt(size)
        if piece.startswith("**"):
            run.bold = True
        elif piece.startswith("*"):
            run.italic = True
        elif piece.startswith("`"):
            run.font.name = "Consolas"
        elif piece.startswith("["):
            run.bold = True
            run.font.color.rgb = _OPEN_ITEM


def _blocks(text: str) -> list[tuple[str, object]]:
    lines = [re.sub(r"<!--.*?-->", "", line).rstrip() for line in text.splitlines()]
    out, buffer, kind = [], [], None

    def flush():
        nonlocal buffer, kind
        if buffer:
            out.append((kind, buffer))
        buffer, kind = [], None

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush()
        elif stripped.startswith("# "):
            flush()
            out.append(("title", stripped[2:]))
        elif stripped.startswith("## "):
            flush()
            out.append(("heading", stripped[3:]))
        elif set(stripped) == {"-"} and len(stripped) >= 3:
            flush()
            out.append(("rule", None))
        elif stripped.startswith("|"):
            if kind != "table":
                flush()
                kind = "table"
            buffer.append(stripped)
        elif stripped.startswith("- "):
            if kind != "bullets":
                flush()
                kind = "bullets"
            buffer.append(stripped[2:])
        elif kind == "bullets":
            buffer[-1] += " " + stripped
        else:
            if kind != "paragraph":
                flush()
                kind = "paragraph"
            buffer.append(stripped)
    flush()
    return out


def _table(doc, rows: list[str]) -> None:
    """The one thing this needs beyond ``kit.table``: parsing a markdown
    pipe table (alignment from the ``---:``/``---`` spec row, columns sized
    to fit the content) instead of taking Python values directly. The
    header/shading/row-building loop itself is not duplicated here (fix
    round 1, minor finding) -- it is built by ``kit.table``, told to align
    its header cells like the body via ``align_headers``."""
    head = [c.replace("\\|", "|") for c in _split_row(rows[0])]
    cells = [head] + [[c.replace("\\|", "|") for c in _split_row(r, len(head) - 1)] for r in rows[1:]]
    spec, body = cells[1], cells[2:]
    aligns = ["right" if c.endswith(":") and not c.startswith(":") else "left" for c in spec]
    longest = [max(len(r[i]) for r in cells if i < len(r)) or 1 for i in range(len(head))]
    raw_widths = [max(1.5, kit.USABLE * n / sum(longest)) for n in longest]
    squeeze = kit.USABLE / sum(raw_widths)
    widths = [w * squeeze for w in raw_widths]
    grid = kit.table(
        doc,
        [h.replace("**", "") for h in head],
        [[v.replace("**", "") for v in row[:len(head)]] for row in body],
        widths=widths, aligns=aligns, align_headers=True,
    )
    for row in grid.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def render_docx(md_text: str, out: pathlib.Path, case: CaseData) -> None:
    doc = kit.new_doc(f"{case.project.name}, glare narrative, case {case.case}, page")
    for kind, payload in _blocks(md_text):
        if kind == "title":
            doc.add_heading(payload, level=0)
            note = doc.add_paragraph()
            run = note.add_run(f"{case.project.name}. Rendered from glare_narrative.md, the file to edit. "
                               f"Case {case.case}, {case.created.strftime('%d %B %Y')}.")
            run.italic = True
            run.font.size = Pt(9)
        elif kind == "heading":
            doc.add_heading(payload, level=1)
        elif kind == "rule":
            paragraph = doc.add_paragraph()
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "BBBBBB")):
                bottom.set(qn(key), value)
            border.append(bottom)
            paragraph._p.get_or_add_pPr().append(border)
        elif kind == "table":
            _table(doc, payload)
        elif kind == "bullets":
            for item in payload:
                _runs(doc.add_paragraph(style="List Bullet"), item)
        else:
            _runs(doc.add_paragraph(), " ".join(payload))
    with explain_in_use(out):
        doc.save(out)
