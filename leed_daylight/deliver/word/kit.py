# leed_daylight/deliver/word/kit.py
"""The Word layout every report shares: A4 page, Calibri 10 pt, grey table
headers and totals, key/value tables, bullets, captions and notes. Ported
from the author's earlier work (PORTED.md); it reads no results and no paths."""
from __future__ import annotations

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

USABLE = 17.4  # text width of an A4 portrait page with 1.8 cm margins, cm
GREY_HEADER = "D9D9D9"
GREY_TOTAL = "F2F2F2"
_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER}


def shade(cell, hex_colour: str) -> None:
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(element)


def put(cell, text, bold: bool = False, align: str = "left", size: int = 9) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = _ALIGN[align]
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)


def repeat_header(row) -> None:
    """Mark a table row as a header, so Word repeats it on every page."""
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(element)


def table(doc, headers, rows, widths=None, aligns=None, total=None, align_headers=False):
    """Build a bordered table. ``align_headers`` (used by the narrative's
    markdown-table renderer, which reads column alignment from the markdown
    spec row) applies ``aligns`` to the header cells too, instead of the
    usual left-aligned header text."""
    grid = doc.add_table(rows=1, cols=len(headers))
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.LEFT
    grid.autofit = widths is None
    aligns = aligns or ["left"] * len(headers)
    for i, header in enumerate(headers):
        put(grid.rows[0].cells[i], header, bold=True, align=(aligns[i] if align_headers else "left"))
        shade(grid.rows[0].cells[i], GREY_HEADER)
    repeat_header(grid.rows[0])
    for row in rows:
        cells = grid.add_row().cells
        for i, value in enumerate(row):
            put(cells[i], value, align=aligns[i])
    if total:
        cells = grid.add_row().cells
        for i, value in enumerate(total):
            put(cells[i], value, bold=True, align=aligns[i])
            shade(cells[i], GREY_TOTAL)
    if widths:
        for row in grid.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    return grid


def kv_table(doc, pairs, key_width: float = 6.2, value_width: float = 11.2):
    grid = doc.add_table(rows=0, cols=2)
    grid.style = "Table Grid"
    for key, value in pairs:
        cells = grid.add_row().cells
        put(cells[0], key, bold=True)
        shade(cells[0], GREY_HEADER)
        put(cells[1], value)
        cells[0].width, cells[1].width = Cm(key_width), Cm(value_width)
    return grid


def bullets(doc, lines) -> None:
    for line in lines:
        paragraph = doc.add_paragraph(line, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)


def caption(doc, text: str, before: int = 0) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(9)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def note(doc, lead: str, body: str):
    paragraph = doc.add_paragraph()
    paragraph.add_run(lead).bold = True
    paragraph.add_run(body)
    return paragraph


def new_doc(footer_text: str):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)  # F16: A4, 1.8 cm margins on every side
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(footer_text + "    ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)
    return doc
